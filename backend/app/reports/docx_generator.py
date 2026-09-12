"""
backend/app/reports/docx_generator.py

Production Microsoft Word (.docx) Inspection Report Generator for LABELGUARD.
Uses python-docx to generate fully editable, inspection-ready legal metrology reports
from the canonical report payload (build_inspection_report_data).
"""

from datetime import datetime
import os
from pathlib import Path
from typing import Any

import docx
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from sqlalchemy.orm import Session

from app.reports.report_data import build_inspection_report_data


REPORTS_DIR = Path(__file__).resolve().parents[2] / "storage" / "reports"

# Document dimensions
PAGE_WIDTH_INCHES = 8.27  # A4 width
PAGE_HEIGHT_INCHES = 11.69  # A4 height
MARGIN_INCHES = 0.5
CONTENT_WIDTH_INCHES = PAGE_WIDTH_INCHES - (2 * MARGIN_INCHES)  # 7.27 inches

# Color constants
COLOR_PRIMARY_NAVY = "1A365D"
COLOR_SECONDARY_BLUE = "2B6CB0"
COLOR_DARK_TEXT = "2D3748"
COLOR_MUTED_GRAY = "718096"
COLOR_LIGHT_BG = "F8FAFC"
COLOR_BORDER = "CBD5E0"
COLOR_PASS_BG = "C6F6D5"
COLOR_PASS_FG = "1C4532"
COLOR_FAIL_BG = "FED7D7"
COLOR_FAIL_FG = "742A2A"
COLOR_REVIEW_BG = "FEEBC8"
COLOR_REVIEW_FG = "744210"
COLOR_NEUTRAL_BG = "EDF2F7"


def _set_cell_background(cell: Any, hex_color: str) -> None:
    """Sets background fill color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(
    cell: Any,
    top: int = 80,
    bottom: int = 80,
    left: int = 120,
    right: int = 120,
) -> None:
    """Sets internal padding (in twips, 20 twips = 1 pt) for a cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for m_name, m_val in [
        ("top", top),
        ("bottom", bottom),
        ("left", left),
        ("right", right),
    ]:
        node = OxmlElement(f"w:{m_name}")
        node.set(qn("w:w"), str(m_val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def _set_table_borders(table: Any, hex_color: str = COLOR_BORDER) -> None:
    """Applies subtle borders to a table."""
    tblPr = table._tbl.tblPr
    borders_xml = f"""
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:left w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>
    </w:tblBorders>
    """
    tblPr.append(parse_xml(borders_xml))


def _set_table_header_repeat(table: Any) -> None:
    """Configures the first row of a table to repeat as header across page breaks."""
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:tblHeader"))


def _set_row_cant_split(row: Any) -> None:
    """Prevents a table row from splitting across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _get_status_badge_info(status: str) -> tuple[str, str, str]:
    """Returns (fg_color, bg_color, display_text) for a status value."""
    st = str(status).strip().lower()
    if st in ("compliant", "pass"):
        return COLOR_PASS_FG, COLOR_PASS_BG, "COMPLIANT / PASS"
    if st in ("non_compliant", "fail", "failed"):
        return COLOR_FAIL_FG, COLOR_FAIL_BG, "NON-COMPLIANT / FAIL"
    if st in ("review", "warning"):
        return COLOR_REVIEW_FG, COLOR_REVIEW_BG, "REVIEW REQUIRED"
    return COLOR_DARK_TEXT, COLOR_NEUTRAL_BG, st.upper() if st else "PENDING"


def _add_section_heading(doc: docx.Document, title: str) -> None:
    """Adds a standardized section heading."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)


def generate_inspection_docx(
    db: Session,
    inspection_id: int,
) -> Path:
    """
    Generates an editable, professional Legal Metrology Word (.docx) inspection report
    using python-docx and saves it to storage/reports/<reference_number>_inspection_report.docx.
    """
    report_data = build_inspection_report_data(
        db=db,
        inspection_id=inspection_id,
        report_format="docx",
    )

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    ref_num = report_data["inspection"]["reference_number"]
    docx_filename = f"{ref_num}_inspection_report.docx"
    docx_path = REPORTS_DIR / docx_filename

    # Initialize Document
    doc = docx.Document()

    # Configure A4 Page Setup and Margins
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH_INCHES)
    section.page_height = Inches(PAGE_HEIGHT_INCHES)
    section.top_margin = Inches(MARGIN_INCHES)
    section.bottom_margin = Inches(MARGIN_INCHES)
    section.left_margin = Inches(MARGIN_INCHES)
    section.right_margin = Inches(MARGIN_INCHES)

    # Configure Running Header & Footer
    header_para = section.header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header_para.add_run(
        f"LABELGUARD — Legal Metrology Inspection Report | Ref: {ref_num}"
    )
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8)
    hrun.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    frun = footer_para.add_run(
        "CONFIDENTIAL — Authorized Regulatory Inspection Document"
    )
    frun.font.name = "Calibri"
    frun.font.size = Pt(8)
    frun.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # ---------------------------------------------------------
    # 1. COVER / HEADER BANNER
    # ---------------------------------------------------------
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False

    col_w_left = Inches(CONTENT_WIDTH_INCHES - 2.2)
    col_w_right = Inches(2.2)

    cell_left = header_table.cell(0, 0)
    cell_right = header_table.cell(0, 1)
    cell_left.width = col_w_left
    cell_right.width = col_w_right

    # Left header info
    p_brand = cell_left.paragraphs[0]
    p_brand.paragraph_format.space_after = Pt(2)
    r_brand = p_brand.add_run("LABELGUARD")
    r_brand.font.name = "Calibri"
    r_brand.font.size = Pt(20)
    r_brand.font.bold = True
    r_brand.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_sub = cell_left.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(4)
    r_sub = p_sub.add_run("AI-Assisted Legal Metrology Inspection Report")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    p_meta = cell_left.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(0)
    r_meta = p_meta.add_run(
        f"Reference: {ref_num}   |   Inspection ID: {report_data['inspection']['inspection_id']}\n"
        f"Report Generated: {report_data['metadata']['report_generated_at']}"
    )
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(8.5)
    r_meta.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # Right header badge
    st_fg, st_bg, st_label = _get_status_badge_info(
        report_data["inspection"]["compliance_status"]
    )
    p_badge = cell_right.paragraphs[0]
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_before = Pt(8)
    r_badge = p_badge.add_run(st_label)
    r_badge.font.name = "Calibri"
    r_badge.font.size = Pt(11)
    r_badge.font.bold = True
    fg_rgb = RGBColor(
        int(st_fg[0:2], 16), int(st_fg[2:4], 16), int(st_fg[4:6], 16)
    )
    r_badge.font.color.rgb = fg_rgb
    _set_cell_background(cell_right, st_bg)
    _set_cell_margins(cell_right, top=140, bottom=140, left=100, right=100)

    _set_cell_margins(cell_left, top=40, bottom=40, left=40, right=40)
    _set_table_borders(header_table, COLOR_BORDER)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(4)
    p_spacer.paragraph_format.space_after = Pt(4)

    # ---------------------------------------------------------
    # 2. INSPECTION DETAILS & LOCATION EVIDENCE (SIDE-BY-SIDE)
    # ---------------------------------------------------------
    _add_section_heading(doc, "1. INSPECTION PROFILE & LOCATION EVIDENCE")

    details_table = doc.add_table(rows=6, cols=4)
    details_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(details_table, COLOR_BORDER)

    w_lbl = Inches(1.3)
    w_val = Inches((CONTENT_WIDTH_INCHES - 2.6) / 2.0)
    for row in details_table.rows:
        row.cells[0].width = w_lbl
        row.cells[1].width = w_val
        row.cells[2].width = w_lbl
        row.cells[3].width = w_val
        _set_row_cant_split(row)

    insp = report_data["inspection"]
    inspector = report_data["inspector"]
    loc = report_data["location"]

    details_matrix = [
        (
            "Inspector:",
            f"{inspector.get('full_name')} ({inspector.get('role')})",
            "GPS Coordinates:",
            (
                f"{loc['latitude']:.6f}, {loc['longitude']:.6f}"
                if loc.get("is_available")
                else "Location unavailable"
            ),
        ),
        (
            "Inspector Email:",
            str(inspector.get("email") or "-"),
            "GPS Accuracy:",
            (
                f"±{loc['location_accuracy_m']} m"
                if loc.get("location_accuracy_m") is not None
                else "-"
            ),
        ),
        (
            "Inspection Status:",
            f"{str(insp.get('status')).upper()} / {str(insp.get('compliance_status')).upper()}",
            "Location Source:",
            str(loc.get("location_source") or "-"),
        ),
        (
            "Created At:",
            str(insp.get("created_at") or "-"),
            "Location Time:",
            str(loc.get("location_captured_at") or "-"),
        ),
        (
            "Scan Started:",
            str(insp.get("scan_started_at") or "-"),
            "Verification:",
            (
                "Geo-tagged Valid"
                if loc.get("is_available")
                else "Not Available"
            ),
        ),
        (
            "Scan Completed:",
            str(insp.get("scan_completed_at") or "-"),
            "Audit State:",
            "Authenticated Inspector Session",
        ),
    ]

    for row_idx, (l1, v1, l2, v2) in enumerate(details_matrix):
        row = details_table.rows[row_idx]
        for col_idx, text, is_label in [
            (0, l1, True),
            (1, v1, False),
            (2, l2, True),
            (3, v2, False),
        ]:
            cell = row.cells[col_idx]
            _set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            if is_label:
                _set_cell_background(cell, COLOR_LIGHT_BG)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.bold = is_label
            r.font.color.rgb = (
                RGBColor(0x1A, 0x36, 0x5D)
                if is_label
                else RGBColor(0x2D, 0x37, 0x48)
            )

    # ---------------------------------------------------------
    # 3. PRODUCT DETAILS
    # ---------------------------------------------------------
    _add_section_heading(doc, "2. PRODUCT DETAILS")

    prod = report_data.get("product")
    if prod is not None:
        prod_table = doc.add_table(rows=3, cols=4)
        prod_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(prod_table, COLOR_BORDER)

        for row in prod_table.rows:
            row.cells[0].width = w_lbl
            row.cells[1].width = w_val
            row.cells[2].width = w_lbl
            row.cells[3].width = w_val
            _set_row_cant_split(row)

        prod_matrix = [
            (
                "Product Name:",
                str(prod.get("product_name") or "-"),
                "Brand Name:",
                str(prod.get("brand_name") or "-"),
            ),
            (
                "Category:",
                str(prod.get("category") or "-"),
                "Package Type:",
                str(prod.get("package_type") or "-"),
            ),
            (
                "Manufacturer:",
                str(prod.get("manufacturer_name") or "-"),
                "Address:",
                str(prod.get("manufacturer_address") or "-"),
            ),
        ]

        for row_idx, (l1, v1, l2, v2) in enumerate(prod_matrix):
            row = prod_table.rows[row_idx]
            for col_idx, text, is_label in [
                (0, l1, True),
                (1, v1, False),
                (2, l2, True),
                (3, v2, False),
            ]:
                cell = row.cells[col_idx]
                _set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
                if is_label:
                    _set_cell_background(cell, COLOR_LIGHT_BG)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(text)
                r.font.name = "Calibri"
                r.font.size = Pt(8.5)
                r.font.bold = is_label
                r.font.color.rgb = (
                    RGBColor(0x1A, 0x36, 0x5D)
                    if is_label
                    else RGBColor(0x2D, 0x37, 0x48)
                )
    else:
        p_unlinked = doc.add_paragraph()
        p_unlinked.paragraph_format.space_after = Pt(2)
        r_un = p_unlinked.add_run(
            "No registered catalog product associated with this inspection (Unlinked Package Scan)."
        )
        r_un.font.name = "Calibri"
        r_un.font.size = Pt(8.5)
        r_un.font.italic = True

    # ---------------------------------------------------------
    # 4. COMPLIANCE SUMMARY METRICS
    # ---------------------------------------------------------
    _add_section_heading(doc, "3. COMPLIANCE EVALUATION SUMMARY")

    summ = report_data["evidence_summary"]
    metrics_table = doc.add_table(rows=2, cols=7)
    metrics_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(metrics_table, COLOR_BORDER)
    _set_table_header_repeat(metrics_table)

    col_w_m = Inches(CONTENT_WIDTH_INCHES / 7.0)
    for row in metrics_table.rows:
        for c in row.cells:
            c.width = col_w_m
        _set_row_cant_split(row)

    headers_metrics = [
        "Total Rules",
        "Passed",
        "Failed",
        "Review",
        "N/A",
        "Declarations",
        "Violations",
    ]
    vals_metrics = [
        str(summ.get("total_rules", 0)),
        str(summ.get("passed_rules", 0)),
        str(summ.get("failed_rules", 0)),
        str(summ.get("reviewed_rules", 0)),
        str(summ.get("not_applicable_rules", 0)),
        str(summ.get("declarations_count", 0)),
        str(summ.get("violations_count", 0)),
    ]

    for col_idx, (hdr, val) in enumerate(zip(headers_metrics, vals_metrics)):
        # Header cell
        c_hdr = metrics_table.cell(0, col_idx)
        _set_cell_background(c_hdr, COLOR_PRIMARY_NAVY)
        _set_cell_margins(c_hdr, top=80, bottom=80, left=40, right=40)
        p_h = c_hdr.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h.paragraph_format.space_after = Pt(0)
        r_h = p_h.add_run(hdr)
        r_h.font.name = "Calibri"
        r_h.font.size = Pt(8)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Value cell
        c_val = metrics_table.cell(1, col_idx)
        _set_cell_background(c_val, COLOR_LIGHT_BG)
        _set_cell_margins(c_val, top=80, bottom=80, left=40, right=40)
        p_v = c_val.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_v.paragraph_format.space_after = Pt(0)
        r_v = p_v.add_run(val)
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(10)
        r_v.font.bold = True
        if hdr == "Passed":
            r_v.font.color.rgb = RGBColor(0x1C, 0x45, 0x32)
        elif hdr == "Failed" and int(val) > 0:
            r_v.font.color.rgb = RGBColor(0x74, 0x2A, 0x2A)
        elif hdr == "Review" and int(val) > 0:
            r_v.font.color.rgb = RGBColor(0x74, 0x42, 0x10)
        else:
            r_v.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    # ---------------------------------------------------------
    # 5. ACTUAL SCANNED IMAGES (ORIGINAL EVIDENCE)
    # ---------------------------------------------------------
    images = report_data.get("images", [])
    if images:
        _add_section_heading(
            doc, "4. ACTUAL SCANNED IMAGES (ORIGINAL EVIDENCE)"
        )

        for idx, img_item in enumerate(images):
            abs_path = img_item.get("absolute_path")
            file_name = img_item.get("file_name", "image.png")
            panel = str(img_item.get("image_type") or "Scanned").capitalize()

            img_box_table = doc.add_table(rows=1, cols=2)
            img_box_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            _set_table_borders(img_box_table, COLOR_BORDER)

            c_img = img_box_table.cell(0, 0)
            c_meta = img_box_table.cell(0, 1)
            c_img.width = Inches(3.2)
            c_meta.width = Inches(CONTENT_WIDTH_INCHES - 3.2)
            _set_cell_margins(c_img, top=100, bottom=100, left=100, right=100)
            _set_cell_margins(
                c_meta, top=100, bottom=100, left=100, right=100
            )

            # Insert Image or Placeholder
            p_img = c_img.paragraphs[0]
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_embedded = False

            if abs_path and os.path.exists(abs_path):
                try:
                    with PILImage.open(abs_path) as pil_img:
                        orig_w, orig_h = pil_img.size

                    aspect = orig_w / float(orig_h)
                    max_w = 2.8
                    max_h = 2.0

                    if orig_w > orig_h:
                        w_in = min(max_w, orig_w / 100.0)
                        h_in = w_in / aspect
                        if h_in > max_h:
                            h_in = max_h
                            w_in = h_in * aspect
                    else:
                        h_in = min(max_h, orig_h / 100.0)
                        w_in = h_in * aspect
                        if w_in > max_w:
                            w_in = max_w
                            h_in = w_in / aspect

                    run_pic = p_img.add_run()
                    run_pic.add_picture(
                        abs_path, width=Inches(w_in), height=Inches(h_in)
                    )
                    img_embedded = True
                except Exception:
                    img_embedded = False

            if not img_embedded:
                r_ph = p_img.add_run(
                    f"[{panel} Panel]\n{file_name}\n[Image unavailable]"
                )
                r_ph.font.name = "Calibri"
                r_ph.font.size = Pt(8.5)
                r_ph.font.italic = True
                r_ph.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

            # Image Metadata text
            p_meta = c_meta.paragraphs[0]
            p_meta.paragraph_format.space_after = Pt(2)
            r_head = p_meta.add_run(f"Image {idx + 1}: {panel} Panel\n")
            r_head.font.name = "Calibri"
            r_head.font.size = Pt(9.5)
            r_head.font.bold = True
            r_head.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

            r_m = p_meta.add_run(
                f"File Name: {file_name}\n"
                f"Image ID: {img_item.get('id')}\n"
                f"MIME Type: {img_item.get('mime_type')}\n"
                f"File Size: {img_item.get('file_size', 0):,} bytes\n"
                f"Content Hash: {str(img_item.get('content_hash'))[:32]}...\n"
                f"Uploaded At: {str(img_item.get('created_at'))[:19]}"
            )
            r_m.font.name = "Calibri"
            r_m.font.size = Pt(8)
            r_m.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)

    # ---------------------------------------------------------
    # 6. EXTRACTED DECLARATIONS TABLE
    # ---------------------------------------------------------
    _add_section_heading(doc, "5. MANDATORY DECLARATIONS EVIDENCE")

    declarations = report_data.get("declarations", [])
    num_decl_rows = len(declarations) + 1 if declarations else 2
    decl_table = doc.add_table(rows=num_decl_rows, cols=5)
    decl_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(decl_table, COLOR_BORDER)
    _set_table_header_repeat(decl_table)

    decl_widths = [
        Inches(1.8),
        Inches(2.2),
        Inches(1.8),
        Inches(0.7),
        Inches(0.77),
    ]
    for row in decl_table.rows:
        for idx_w, w_val in enumerate(decl_widths):
            row.cells[idx_w].width = w_val
        _set_row_cant_split(row)

    decl_headers = [
        "Declaration Field",
        "Extracted Value (Raw OCR)",
        "Normalized Value",
        "Present",
        "Confidence",
    ]
    for col_idx, hdr in enumerate(decl_headers):
        c_hdr = decl_table.cell(0, col_idx)
        _set_cell_background(c_hdr, COLOR_PRIMARY_NAVY)
        _set_cell_margins(c_hdr, top=80, bottom=80, left=60, right=60)
        p = c_hdr.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(hdr)
        r.font.name = "Calibri"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if declarations:
        for row_idx, d in enumerate(declarations, start=1):
            row = decl_table.rows[row_idx]
            conf_str = (
                f"{d['confidence']:.1f}%"
                if d.get("confidence") is not None
                else "-"
            )
            pres_str = "YES" if d.get("is_present") else "NO"

            row_vals = [
                (str(d.get("field_name") or "-"), True),
                (str(d.get("extracted_value") or "-"), False),
                (str(d.get("normalized_value") or "-"), False),
                (pres_str, False),
                (conf_str, False),
            ]

            bg_color = COLOR_LIGHT_BG if row_idx % 2 == 0 else "FFFFFF"
            for col_idx, (text, is_bold) in enumerate(row_vals):
                cell = row.cells[col_idx]
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(text)
                r.font.name = "Calibri"
                r.font.size = Pt(8)
                r.font.bold = is_bold
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    else:
        row = decl_table.rows[1]
        for col_idx in range(5):
            cell = row.cells[col_idx]
            _set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(
                "No declarations extracted." if col_idx == 0 else "-"
            )
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.italic = True

    # ---------------------------------------------------------
    # 7. RULE COMPLIANCE MATRIX (DYNAMIC DDL)
    # ---------------------------------------------------------
    _add_section_heading(
        doc, "6. STATUTORY COMPLIANCE RULE EVALUATION MATRIX"
    )

    rules = report_data.get("rules", [])
    num_rule_rows = len(rules) + 1 if rules else 2
    rule_table = doc.add_table(rows=num_rule_rows, cols=4)
    rule_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(rule_table, COLOR_BORDER)
    _set_table_header_repeat(rule_table)

    rule_widths = [Inches(1.6), Inches(2.8), Inches(1.2), Inches(1.67)]
    for row in rule_table.rows:
        for idx_w, w_val in enumerate(rule_widths):
            row.cells[idx_w].width = w_val
        _set_row_cant_split(row)

    rule_headers = [
        "Rule / Statute",
        "Requirement / Title",
        "Result",
        "Evaluated Checks",
    ]
    for col_idx, hdr in enumerate(rule_headers):
        c_hdr = rule_table.cell(0, col_idx)
        _set_cell_background(c_hdr, COLOR_PRIMARY_NAVY)
        _set_cell_margins(c_hdr, top=80, bottom=80, left=60, right=60)
        p = c_hdr.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(hdr)
        r.font.name = "Calibri"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for row_idx, r_item in enumerate(rules, start=1):
        row = rule_table.rows[row_idx]
        r_fg, r_bg, r_label = _get_status_badge_info(
            r_item.get("result_status")
        )

        checks_text = "\n".join(
            f"• {c.get('field_name')} ({c.get('operator')}): {str(c.get('status')).upper()}"
            for c in r_item.get("checks", [])
        )

        cell_data = [
            f"{r_item.get('rule_code')}\n{r_item.get('rule_number')} (v{r_item.get('version')})",
            f"{r_item.get('title')}\n{r_item.get('requirement')}",
            r_label,
            checks_text or "-",
        ]

        bg_color = COLOR_LIGHT_BG if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(cell_data):
            cell = row.cells[col_idx]
            _set_cell_background(cell, bg_color)
            _set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(8)
            if col_idx == 0:
                run.font.bold = True
            elif col_idx == 2:
                run.font.bold = True
                fg_rgb = RGBColor(
                    int(r_fg[0:2], 16),
                    int(r_fg[2:4], 16),
                    int(r_fg[4:6], 16),
                )
                run.font.color.rgb = fg_rgb

    # ---------------------------------------------------------
    # 8. VIOLATIONS & REVIEW FINDINGS
    # ---------------------------------------------------------
    _add_section_heading(doc, "7. RECORDED VIOLATIONS & REVIEW FINDINGS")

    violations = report_data.get("violations", [])
    if not violations:
        p_noviol = doc.add_paragraph()
        p_noviol.paragraph_format.space_after = Pt(4)
        r_nv = p_noviol.add_run(
            "No violations identified. All scanned label declarations satisfy mandatory legal metrology requirements."
        )
        r_nv.font.name = "Calibri"
        r_nv.font.size = Pt(9)
        r_nv.font.bold = True
        r_nv.font.color.rgb = RGBColor(0x1C, 0x45, 0x32)
    else:
        num_v_rows = len(violations) + 1
        viol_table = doc.add_table(rows=num_v_rows, cols=5)
        viol_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(viol_table, "E53E3E")
        _set_table_header_repeat(viol_table)

        viol_widths = [
            Inches(1.4),
            Inches(1.0),
            Inches(1.1),
            Inches(2.0),
            Inches(1.77),
        ]
        for row in viol_table.rows:
            for idx_w, w_val in enumerate(viol_widths):
                row.cells[idx_w].width = w_val
            _set_row_cant_split(row)

        viol_headers = [
            "Field / Rule",
            "Severity",
            "Status",
            "Violation Message",
            "Detected vs Expected",
        ]
        for col_idx, hdr in enumerate(viol_headers):
            c_hdr = viol_table.cell(0, col_idx)
            _set_cell_background(c_hdr, "742A2A")
            _set_cell_margins(c_hdr, top=80, bottom=80, left=60, right=60)
            p = c_hdr.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(hdr)
            r.font.name = "Calibri"
            r.font.size = Pt(8)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        for row_idx, v in enumerate(violations, start=1):
            row = viol_table.rows[row_idx]
            v_fg, v_bg, v_status = _get_status_badge_info(v.get("status"))
            detected_str = f"Det: {v.get('detected_value') or '-'}\nExp: {v.get('expected_value') or '-'}"

            row_data = [
                str(v.get("field_name") or "-"),
                str(v.get("severity") or "-").upper(),
                v_status,
                str(v.get("message") or "-"),
                detected_str,
            ]

            bg_color = "FFF5F5" if row_idx % 2 == 0 else "FFFFFF"
            for col_idx, text in enumerate(row_data):
                cell = row.cells[col_idx]
                _set_cell_background(cell, bg_color)
                _set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(text)
                run.font.name = "Calibri"
                run.font.size = Pt(8)
                if col_idx in (0, 1):
                    run.font.bold = True
                elif col_idx == 2:
                    run.font.bold = True
                    fg_rgb = RGBColor(
                        int(v_fg[0:2], 16),
                        int(v_fg[2:4], 16),
                        int(v_fg[4:6], 16),
                    )
                    run.font.color.rgb = fg_rgb

    # ---------------------------------------------------------
    # 9. REGULATORY TRACEABILITY & AUDIT PROVENANCE
    # ---------------------------------------------------------
    _add_section_heading(doc, "8. REGULATORY TRACEABILITY & AUDIT PROVENANCE")

    reg = report_data.get("regulation") or {}
    trace_table = doc.add_table(rows=2, cols=4)
    trace_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(trace_table, COLOR_BORDER)

    for row in trace_table.rows:
        row.cells[0].width = w_lbl
        row.cells[1].width = w_val
        row.cells[2].width = w_lbl
        row.cells[3].width = w_val
        _set_row_cant_split(row)

    trace_matrix = [
        (
            "Enacting Regulation:",
            f"{reg.get('name')} ({reg.get('code')})",
            "Jurisdiction / Authority:",
            f"{reg.get('jurisdiction')} — {reg.get('authority')}",
        ),
        (
            "Generator Engine:",
            f"{report_data['metadata'].get('generator_version')} (python-docx 1.2.0)",
            "Report Generated At:",
            str(report_data["metadata"].get("report_generated_at")),
        ),
    ]

    for row_idx, (l1, v1, l2, v2) in enumerate(trace_matrix):
        row = trace_table.rows[row_idx]
        for col_idx, text, is_label in [
            (0, l1, True),
            (1, v1, False),
            (2, l2, True),
            (3, v2, False),
        ]:
            cell = row.cells[col_idx]
            _set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            if is_label:
                _set_cell_background(cell, COLOR_LIGHT_BG)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.name = "Calibri"
            r.font.size = Pt(8.5)
            r.font.bold = is_label
            r.font.color.rgb = (
                RGBColor(0x1A, 0x36, 0x5D)
                if is_label
                else RGBColor(0x2D, 0x37, 0x48)
            )

    # ---------------------------------------------------------
    # 10. STATUTORY DISCLAIMER
    # ---------------------------------------------------------
    p_disc_head = doc.add_paragraph()
    p_disc_head.paragraph_format.space_before = Pt(8)
    p_disc_head.paragraph_format.space_after = Pt(2)
    r_dh = p_disc_head.add_run("STATUTORY NOTICE & DISCLAIMER:")
    r_dh.font.name = "Calibri"
    r_dh.font.size = Pt(8)
    r_dh.font.bold = True
    r_dh.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    p_disc = doc.add_paragraph()
    p_disc.paragraph_format.space_after = Pt(8)
    r_d = p_disc.add_run(
        "This inspection record is generated by LABELGUARD AI as an automated decision-support finding "
        "under the Legal Metrology (Packaged Commodities) Rules, 2011. This document constitutes preliminary technical "
        "analysis. Final statutory enforcement, compounding, or adjudication remains strictly within the discretionary "
        "authority of the designated Legal Metrology Officer."
    )
    r_d.font.name = "Calibri"
    r_d.font.size = Pt(8)
    r_d.font.italic = True
    r_d.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    # ---------------------------------------------------------
    # 11. OFFICER REVIEW & DETERMINATION (EDITABLE FIELDS)
    # ---------------------------------------------------------
    _add_section_heading(doc, "9. OFFICER REVIEW & FINAL DETERMINATION")

    officer_table = doc.add_table(rows=4, cols=2)
    officer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(officer_table, COLOR_BORDER)

    officer_table.rows[0].cells[0].width = Inches(2.2)
    officer_table.rows[0].cells[1].width = Inches(CONTENT_WIDTH_INCHES - 2.2)

    officer_matrix = [
        (
            "Officer Remarks / Observations:",
            "\n\n____________________________________________________________________________________",
        ),
        (
            "Final Statutory Determination:",
            "[   ] APPROVED / COMPLIANT\n[   ] VIOLATION NOTICE ISSUED\n[   ] RE-INSPECTION / SEIZURE ORDERED",
        ),
        (
            "Authorized Officer Name & Title:",
            "____________________________________________________________________________________",
        ),
        (
            "Signature & Date:",
            "Signature: ___________________________          Date: ________________________",
        ),
    ]

    for row_idx, (label, val) in enumerate(officer_matrix):
        row = officer_table.rows[row_idx]
        _set_row_cant_split(row)

        c_lbl = row.cells[0]
        c_val = row.cells[1]
        c_lbl.width = Inches(2.2)
        c_val.width = Inches(CONTENT_WIDTH_INCHES - 2.2)

        _set_cell_background(c_lbl, COLOR_LIGHT_BG)
        _set_cell_margins(c_lbl, top=80, bottom=80, left=80, right=80)
        _set_cell_margins(c_val, top=80, bottom=80, left=80, right=80)

        p_l = c_lbl.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(0)
        r_l = p_l.add_run(label)
        r_l.font.name = "Calibri"
        r_l.font.size = Pt(8.5)
        r_l.font.bold = True
        r_l.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

        p_v = c_val.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(0)
        r_v = p_v.add_run(val)
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(8.5)
        r_v.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Save to storage/reports/<reference_number>_inspection_report.docx
    doc.save(str(docx_path))

    return docx_path
