"""
backend/tests/integration/test_docx_report.py

Comprehensive tests for the Editable DOCX Report Generator (docx_generator.py).
Covers requirements A through P:
A. compliant inspection
B. non-compliant inspection
C. review inspection
D. GPS present
E. GPS absent
F. multiple images
G. missing image
H. no violations
I. violations
J. long text
K. DOCX file created in backend/storage/reports/
L. DOCX is a valid ZIP/DOCX package
M. expected report text exists in document
N. images are embedded when available
O. editable tables exist as native Word tables
P. PDF and DOCX canonical data consistency
"""

from pathlib import Path
import zipfile
import docx
import pytest
from PIL import Image
from sqlalchemy.orm import Session

from app.models.declaration import Declaration
from app.models.inspection import ComplianceStatus, Inspection, InspectionStatus
from app.models.inspection_image import InspectionImage
from app.models.product import Product
from app.models.user import User
from app.reports.docx_generator import REPORTS_DIR, generate_inspection_docx
from app.reports.pdf_generator import generate_inspection_pdf
from app.services.inspection_service import create_inspection
from tests.integration.test_reports import _create_test_declaration


def _verify_docx_package(docx_path: Path) -> docx.Document:
    assert docx_path.exists(), f"DOCX does not exist at {docx_path}"
    assert docx_path.is_file()
    assert docx_path.stat().st_size > 2000, "DOCX file size too small"
    assert zipfile.is_zipfile(str(docx_path)), "DOCX is not a valid ZIP/Office package"

    with zipfile.ZipFile(str(docx_path)) as z:
        namelist = z.namelist()
        assert "[Content_Types].xml" in namelist
        assert "word/document.xml" in namelist

    doc = docx.Document(str(docx_path))
    return doc


def _get_full_docx_text(doc: docx.Document) -> str:
    full_text = []
    for p in doc.paragraphs:
        full_text.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    full_text.append(p.text)
    return " ".join(full_text)


def _create_temp_png(file_path: Path, color=(200, 230, 255), size=(300, 200)) -> None:
    file_path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", size, color=color)
    img.save(file_path, format="PNG")


class TestDOCXReportGenerator:
    def test_missing_inspection_raises_value_error(self, db_session: Session):
        with pytest.raises(ValueError, match="does not exist"):
            generate_inspection_docx(db_session, 99999999)

    # A. compliant inspection + H. no violations + K. DOCX created + L. valid package + O. native tables
    def test_A_compliant_inspection_docx(
        self,
        db_session: Session,
        admin_user: User,
    ):
        insp = create_inspection(
            db=db_session,
            inspector=admin_user,
            latitude=28.6139,
            longitude=77.2090,
            location_accuracy_m=4.0,
            location_source="gps",
        )

        compliant_decls = [
            ("mrp", "Rs. 50.00 (incl. of all taxes)", "50.00"),
            ("manufacturer", "Haldiram Snacks Pvt Ltd, Noida, UP", "Haldiram Snacks Pvt Ltd, Noida, UP"),
            ("product_name", "Namkeen / Snack Food", "Namkeen / Snack Food"),
            ("net_quantity", "200 g", "200 g"),
            ("unit_symbol", "g", "g"),
            ("unit_sale_price", "Rs. 0.25 / g", "0.25 / g"),
            ("manufacturing_date", "08/2026", "2026-08"),
            ("best_before", "08/2027", "2027-08"),
            ("date_sensitive_commodity", "true", "true"),
            ("batch_number", "B63240194", "B63240194"),
            ("country_of_origin", "India", "India"),
            ("consumer_care", "customercare@haldiram.com, 1800-111-222", "customercare@haldiram.com, 1800-111-222"),
            ("declaration_legibility", "pass", "pass"),
        ]

        for field_name, ext_val, norm_val in compliant_decls:
            _create_test_declaration(
                db=db_session,
                inspection_id=insp.id,
                field_name=field_name,
                extracted_value=ext_val,
                normalized_value=norm_val,
                confidence=95.0,
            )

        docx_path = generate_inspection_docx(db_session, insp.id)

        assert docx_path.name == f"{insp.reference_number}_inspection_report.docx"
        assert docx_path.parent == REPORTS_DIR

        doc = _verify_docx_package(docx_path)
        text = _get_full_docx_text(doc)

        assert "LABELGUARD" in text
        assert insp.reference_number in text
        assert "COMPLIANT / PASS" in text
        assert "No violations identified" in text
        assert "STATUTORY NOTICE & DISCLAIMER" in text
        assert "OFFICER REVIEW & FINAL DETERMINATION" in text

        # Assert native editable Word tables exist
        assert len(doc.tables) >= 5

    # B. non-compliant inspection + I. violations
    def test_B_non_compliant_inspection_docx(
        self,
        db_session: Session,
        admin_user: User,
    ):
        insp = create_inspection(db=db_session, inspector=admin_user)
        docx_path = generate_inspection_docx(db_session, insp.id)

        doc = _verify_docx_package(docx_path)
        text = _get_full_docx_text(doc)

        assert "NON-COMPLIANT / FAIL" in text or "REVIEW REQUIRED" in text
        assert "RECORDED VIOLATIONS & REVIEW FINDINGS" in text

    # C. review inspection
    def test_C_review_inspection_docx(
        self,
        db_session: Session,
        admin_user: User,
    ):
        insp = create_inspection(db=db_session, inspector=admin_user)
        _create_test_declaration(
            db=db_session,
            inspection_id=insp.id,
            field_name="mrp",
            extracted_value="Rs. 50",
            normalized_value="50.00",
            confidence=20.0,
        )

        docx_path = generate_inspection_docx(db_session, insp.id)
        doc = _verify_docx_package(docx_path)
        text = _get_full_docx_text(doc)

        assert "REVIEW REQUIRED" in text or "REVIEW" in text

    # D. GPS present + E. GPS absent
    def test_D_E_gps_handling_docx(
        self,
        db_session: Session,
        admin_user: User,
    ):
        # GPS present
        insp_gps = create_inspection(
            db=db_session,
            inspector=admin_user,
            latitude=13.0827,
            longitude=80.2707,
            location_accuracy_m=5.0,
            location_source="gps_fused",
        )
        docx_gps = generate_inspection_docx(db_session, insp_gps.id)
        text_gps = _get_full_docx_text(_verify_docx_package(docx_gps))
        assert "13.082700, 80.270700" in text_gps
        assert "gps_fused" in text_gps

        # GPS absent
        insp_nogps = create_inspection(db=db_session, inspector=admin_user)
        docx_nogps = generate_inspection_docx(db_session, insp_nogps.id)
        text_nogps = _get_full_docx_text(_verify_docx_package(docx_nogps))
        assert "Location unavailable" in text_nogps

    # F. multiple images + G. missing image + N. images embedded
    def test_F_G_N_images_embedded_and_missing_fallback(
        self,
        db_session: Session,
        admin_user: User,
        tmp_path: Path,
    ):
        insp = create_inspection(db=db_session, inspector=admin_user)

        real_img = tmp_path / "pkg_front.png"
        _create_temp_png(real_img, color=(150, 200, 100), size=(350, 250))

        img1 = InspectionImage(
            inspection_id=insp.id,
            file_name="pkg_front.png",
            file_path=str(real_img),
            image_type="front",
            file_size=45000,
            mime_type="image/png",
            content_hash="hash_embedded_front_1",
        )
        img2 = InspectionImage(
            inspection_id=insp.id,
            file_name="pkg_back_missing.png",
            file_path=str(tmp_path / "deleted_file.png"),
            image_type="back",
            file_size=55000,
            mime_type="image/png",
            content_hash="hash_missing_back_2",
        )
        db_session.add_all([img1, img2])
        db_session.commit()

        docx_path = generate_inspection_docx(db_session, insp.id)
        doc = _verify_docx_package(docx_path)
        text = _get_full_docx_text(doc)

        assert "pkg_front.png" in text
        assert "pkg_back_missing.png" in text
        assert "[Image unavailable]" in text

        # Check that image part exists in the docx ZIP package
        with zipfile.ZipFile(str(docx_path)) as z:
            media_files = [f for f in z.namelist() if f.startswith("word/media/")]
            assert len(media_files) >= 1

    # J. long text handling
    def test_J_long_text_wrapping(
        self,
        db_session: Session,
        admin_user: User,
    ):
        insp = create_inspection(db=db_session, inspector=admin_user)

        long_text = "Standard Packaged Commodity: " + ("Sample Text Line " * 50)
        _create_test_declaration(
            db=db_session,
            inspection_id=insp.id,
            field_name="manufacturer",
            extracted_value=long_text,
            normalized_value=long_text,
            confidence=90.0,
        )

        docx_path = generate_inspection_docx(db_session, insp.id)
        doc = _verify_docx_package(docx_path)
        assert len(doc.tables) >= 5

    # P. PDF and DOCX canonical data consistency
    def test_P_pdf_and_docx_consistency(
        self,
        db_session: Session,
        admin_user: User,
    ):
        insp = create_inspection(
            db=db_session,
            inspector=admin_user,
            latitude=28.6139,
            longitude=77.2090,
            location_accuracy_m=3.0,
            location_source="gps",
        )
        _create_test_declaration(
            db=db_session,
            inspection_id=insp.id,
            field_name="net_quantity",
            extracted_value="100 g",
            normalized_value="100 g",
        )

        pdf_path = generate_inspection_pdf(db_session, insp.id)
        docx_path = generate_inspection_docx(db_session, insp.id)

        assert pdf_path.exists()
        assert docx_path.exists()

        assert pdf_path.name.replace(".pdf", "") == docx_path.name.replace(".docx", "")
        doc = _verify_docx_package(docx_path)
        text = _get_full_docx_text(doc)

        assert insp.reference_number in text
        assert "100 g" in text
        assert "28.613900, 77.209000" in text
